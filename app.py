import streamlit as st
import joblib
import pandas as pd
import re
import nltk
from nltk.corpus import stopwords
from fpdf import FPDF
import base64
from wordcloud import WordCloud
import matplotlib.pyplot as plt
import PyPDF2
from docx import Document

# Download NLTK data (if not already present)
try:
    nltk.data.find('corpora/stopwords')
except LookupError:
    nltk.download('stopwords')

stop_words = set(stopwords.words('english'))

# Load Model
@st.cache_resource
def load_model():
    try:
        model = joblib.load('model_pipeline.pkl')
        return model
    except FileNotFoundError:
        st.error("Model file 'model_pipeline.pkl' not found. Please run train_model.py first.")
        return None

model = load_model()

def preprocess_text(text):
    if not isinstance(text, str):
        return ""
    text = text.lower()
    text = re.sub(r'[^\w\s]', '', text)
    tokens = text.split()
    tokens = [t for t in tokens if t not in stop_words]
    return " ".join(tokens)

def extract_text(uploaded_file):
    text = ""
    try:
        if uploaded_file.name.endswith('.txt'):
            text = uploaded_file.read().decode("utf-8")
        elif uploaded_file.name.endswith('.pdf'):
            reader = PyPDF2.PdfReader(uploaded_file)
            for page in reader.pages:
                text += page.extract_text() or ""
        elif uploaded_file.name.endswith('.docx'):
            doc = Document(uploaded_file)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
    except Exception as e:
        st.error(f"Error reading file: {e}")
    return text

def create_pdf_report(text, prediction, probability, keywords):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    # Clean text to remove common PDF extraction artifacts
    # Specifically target the '&char&' pattern if it dominates, or odd spacing
    clean_text = text.replace('\x00', '') # Remove null bytes
    if text.count('&') > len(text) / 3:
        clean_text = clean_text.replace('&', '')
    
    # Title
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="Document Classification Report", ln=True, align='C')
    pdf.ln(10)
    
    # Details
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt=f"Predicted Category: {prediction}", ln=True)
    pdf.cell(200, 10, txt=f"Confidence Score: {max(probability):.2f}", ln=True)
    pdf.ln(5)
    
    # Keywords
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(200, 10, txt="Top Keywords:", ln=True)
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, txt=", ".join(keywords))
    pdf.ln(5)
    
    # Content Snippet
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(200, 10, txt="Document Content (Snippet):", ln=True)
    pdf.set_font("Arial", size=10)
    # Filter out characters that might cause latin-1 encode errors in standard FPDF
    safe_text = clean_text[:2000].encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 5, txt=safe_text + "...")
    
    return pdf.output(dest='S').encode('latin-1')

def analyze_text(text, key_suffix=""):
    if not text:
        return

    # Preprocess
    clean_text = preprocess_text(text)
    
    if model:
        # Predict
        prediction = model.predict([clean_text])[0]
        proba = model.predict_proba([clean_text])[0]
        classes = model.classes_
        
        # Display Result
        st.success(f"**Predicted Category:** {prediction}")
        
        # Confidence Scores
        proba_df = pd.DataFrame({
            'Category': classes,
            'Probability': proba
        })
        
        st.write("### Prediction Probabilities")
        st.bar_chart(proba_df.set_index('Category'))

        # Keyword Extraction
        keywords = []
        try:
            vectorizer = model.named_steps['tfidf']
            feature_names = vectorizer.get_feature_names_out()
            tfidf_matrix = vectorizer.transform([clean_text])
            
            # Get top 5 keywords based on TF-IDF score
            sorted_indices = tfidf_matrix.toarray()[0].argsort()[::-1]
            top_n = 5
            top_indices = sorted_indices[:top_n]
            keywords = [feature_names[i] for i in top_indices if tfidf_matrix[0, i] > 0]
            
            st.write("### Top Keywords")
            st.write(", ".join(keywords))
            
            # Word Cloud
            st.write("### Word Cloud")
            try:
                wordcloud = WordCloud(width=800, height=400, background_color='white').generate(text)
                fig, ax = plt.subplots(figsize=(10, 5))
                ax.imshow(wordcloud, interpolation='bilinear')
                ax.axis('off')
                st.pyplot(fig)
            except Exception as wc_err:
                 st.error(f"Error generating word cloud: {wc_err}")
            
            # Highlight keywords in text
            st.write("### Highlighted Content")
            highlighted_text = text[:1000] # Limit to first 1000 chars for display
            for kw in keywords:
                # Case-insensitive replacement
                pattern = re.compile(re.escape(kw), re.IGNORECASE)
                highlighted_text = pattern.sub(f"<mark style='background-color: yellow; color: black;'>{kw}</mark>", highlighted_text)
            
            st.markdown(highlighted_text + ("..." if len(text) > 1000 else ""), unsafe_allow_html=True)
            
        except Exception as e:
            st.error(f"Error extracting keywords: {e}")

        # PDF Report
        st.write("---")
        st.write("### Export Report")
        # Direct download button to avoid nesting issues
        pdf_bytes = create_pdf_report(text, prediction, proba, keywords)
        st.download_button(
            label="Download PDF Report",
            data=pdf_bytes,
            file_name="classification_report.pdf",
            mime="application/pdf",
            key=f"dl_pdf_{key_suffix}"
        )

# UI
st.title("Document Classification System")
st.write("Classify documents into: **Sports, Politics, Technology, Business, or Entertainment**.")

tab1, tab2 = st.tabs(["File Upload", "Live Text Input"])

with tab1:
    uploaded_files = st.file_uploader("Choose files", type=['txt', 'pdf', 'docx'], accept_multiple_files=True)

    if uploaded_files:
        if len(uploaded_files) == 1:
            # Single file mode - Detailed View
            uploaded_file = uploaded_files[0]
            with st.spinner("Extracting text and classifying..."):
                text = extract_text(uploaded_file)
                
                if text:
                    # Show a snippet of the text
                    with st.expander("Show extracted text snippet"):
                        st.write(text[:500] + "..." if len(text) > 500 else text)
                    
                    analyze_text(text, key_suffix="file")
                else:
                    st.warning("Could not extract text from the file.")
                    
        else:
            # Batch Mode - Table View
            results = []
            progress_bar = st.progress(0)
            
            for i, uploaded_file in enumerate(uploaded_files):
                text = extract_text(uploaded_file)
                if text:
                    clean_text = preprocess_text(text)
                    if model:
                        prediction = model.predict([clean_text])[0]
                        proba = model.predict_proba([clean_text])[0]
                        confidence = max(proba)
                        results.append({
                            "Filename": uploaded_file.name,
                            "Predicted Category": prediction,
                            "Confidence": f"{confidence:.2f}"
                        })
                else:
                    results.append({
                        "Filename": uploaded_file.name,
                        "Predicted Category": "Error/Empty",
                        "Confidence": "0.00"
                    })
                progress_bar.progress((i + 1) / len(uploaded_files))
                
            st.success(f"Processed {len(uploaded_files)} files.")
            st.table(pd.DataFrame(results))

with tab2:
    st.write("Paste your text below for classification.")
    user_input = st.text_area("Enter Text", height=250)
    
    # Session state initialization for Live Text
    if "live_text_analyzed" not in st.session_state:
        st.session_state.live_text_analyzed = False
    if "live_text_content" not in st.session_state:
        st.session_state.live_text_content = ""

    if st.button("Classify Text"):
        if user_input.strip():
            st.session_state.live_text_analyzed = True
            st.session_state.live_text_content = user_input.strip()
        else:
            st.warning("Please enter some text to classify.")
    
    # Display results if analyzed
    if st.session_state.live_text_analyzed and st.session_state.live_text_content:
        st.write("---")
        st.write("### Analysis Result")
        analyze_text(st.session_state.live_text_content, key_suffix="text")
